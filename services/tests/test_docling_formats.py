"""Actual pinned converters through a private gRPC server, with portable sources.

PDF model inference is opt-in: HADES_TEST_PDF=/absolute/original.pdf. It uses
CPU, hashes the source before/after, and never writes to a graph or shared service.
"""
import asyncio
import hashlib
import os
from pathlib import Path
import sys

import grpc
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from extraction.server import ExtractionConfig, ExtractionServicer, extraction_pb2 as pb
from persephone.extraction import extraction_pb2_grpc as rpc

TEXT = "The human reviews the source evidence."


def fixtures(root):
    from docx import Document
    from pptx import Presentation
    from openpyxl import Workbook
    from docling_core.types.doc import DoclingDocument, DocItemLabel
    doc = Document()
    doc.add_paragraph(TEXT)
    doc.save(root / "source.docx")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Source"
    slide.placeholders[1].text = TEXT
    deck.save(root / "source.pptx")
    book = Workbook()
    book.active.append(["Evidence"])
    book.active.append([TEXT])
    book.save(root / "source.xlsx")
    docling = DoclingDocument(name="source")
    docling.add_text(label=DocItemLabel.TEXT, text=TEXT)
    docling.save_as_json(root / "source.json")
    for ext, text in {
        "html": f"<!DOCTYPE html><html><body><p>{TEXT}</p></body></html>",
        "md": f"# Source\n\n{TEXT}\n",
        "markdown": f"# Source\n\n{TEXT}\n",
        "txt": TEXT,
        "csv": f"Evidence\n{TEXT}\n",
        "adoc": f"= Source\n\n{TEXT}\n",
        "vtt": f"WEBVTT\n\n00:00:00.000 --> 00:00:01.000\n{TEXT}\n",
    }.items():
        (root / f"source.{ext}").write_text(text)
    return sorted(root.glob("source.*"))


async def check_files(paths):
    service = ExtractionServicer(ExtractionConfig(device="cpu"))
    server = grpc.aio.server()
    rpc.add_ExtractionServiceServicer_to_server(service, server)
    port = server.add_insecure_port("127.0.0.1:0")
    await server.start()
    try:
        async with grpc.aio.insecure_channel(f"127.0.0.1:{port}") as channel:
            stub = rpc.ExtractionServiceStub(channel)
            for path in paths:
                before = hashlib.sha256(path.read_bytes()).hexdigest()
                result = await stub.Extract(pb.ExtractRequest(
                    file_path=str(path), extract_tables=True,
                    extract_equations=True, extract_images=True,
                ), timeout=600)
                assert result.metadata["extractor"] == "docling", path.suffix
                assert result.metadata["docling_version"] == "2.126.0"
                assert result.metadata["docling_core_version"] == "2.96.0"
                assert result.metadata["formula_enrichment"] == "false"
                assert result.full_text.strip(), path.suffix
                assert before == hashlib.sha256(path.read_bytes()).hexdigest()
                if path.suffix != ".pdf":
                    assert TEXT in result.full_text, path.suffix
    finally:
        await server.stop(0)
        await service.close()


def test_real_office_and_text_converters_through_grpc(tmp_path):
    asyncio.run(check_files(fixtures(tmp_path)))


@pytest.mark.skipif(not os.environ.get("HADES_TEST_PDF"), reason="explicit real PDF/model cache required")
def test_real_pdf_through_grpc(tmp_path):
    import shutil
    original = Path(os.environ["HADES_TEST_PDF"])
    digest = hashlib.sha256(original.read_bytes()).hexdigest()
    paths = fixtures(tmp_path)
    copy = tmp_path / "source.pdf"
    shutil.copyfile(original, copy)
    asyncio.run(check_files([*paths, copy]))
    assert hashlib.sha256(original.read_bytes()).hexdigest() == digest


def test_docling_error_reaches_grpc_without_success(tmp_path):
    async def run():
        path = tmp_path / "broken.pdf"
        path.write_bytes(b"not a pdf")
        with pytest.raises(grpc.aio.AioRpcError) as error:
            await check_files([path])
        assert error.value.code() == grpc.StatusCode.INTERNAL
        assert "Docling extraction failed" in error.value.details()
    asyncio.run(run())
