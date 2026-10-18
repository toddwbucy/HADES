#!/usr/bin/env python3
"""Real extraction/daemon/browser test. Invoke only through the isolated DB runner."""
import argparse
import json
import os
from pathlib import Path
import socket
import subprocess
import sys
import time
import urllib.request


def port():
    with socket.socket() as s:
        s.bind(('127.0.0.1', 0))
        return s.getsockname()[1]


def pdf_fixture(path):
    # A small real PDF, no third-party author content and no model-generated argument.
    stream = b'BT /F1 18 Tf 50 740 Td (Guided reading fixture) Tj 0 -40 Td /F1 12 Tf (The human selects a source passage.) Tj 0 -24 Td (The model proposes a note for human review.) Tj ET'
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>', b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
               b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
               b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
               b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream']
    data = bytearray(b'%PDF-1.4\n')
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f'{i} 0 obj\n'.encode() + obj + b'\nendobj\n')
    xref = len(data)
    data.extend(b'xref\n0 6\n0000000000 65535 f \n')
    for offset in offsets[1:]:
        data.extend(f'{offset:010d} 00000 n \n'.encode())
    data.extend(f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode())
    path.write_bytes(data)


def main():
    p = argparse.ArgumentParser(description=__doc__)
    for name in ('root', 'node', 'playwright', 'chromium', 'model-cache'):
        p.add_argument('--'+name, type=Path, required=True)
    a = p.parse_args()
    root = a.root.resolve()
    if os.environ.get('ARANGO_SOCKET') != str(root/'arango.sock') or not (root/'hades.yaml').is_file():
        raise RuntimeError('Use the isolated runner. No ambient database is permitted.')
    repo = Path(__file__).resolve().parents[1]
    target = Path(os.environ['CARGO_TARGET_DIR'])/'debug'
    hades = target/'hades'
    viewer = target/'hades-viewer'
    env = dict(os.environ)
    env.update(HADES_EXTRACTOR_SOCKET=str(root/'extractor.sock'), HADES_EXTRACTOR_DEVICE='cpu',
               HADES_EXTRACTOR_OCR='false', HF_HOME=str(a.model_cache), HF_HUB_OFFLINE='1', OMP_NUM_THREADS='4')
    config = json.loads((root/'hades.yaml').read_text())
    config['extraction']['service']['socket'] = str(root/'extractor.sock')
    (root/'hades.yaml').write_text(json.dumps(config))
    subprocess.run([str(hades),'--db','_system','db','create-database','reading_e2e'],env=env,check=True,capture_output=True)
    pdf_fixture(root/'fixture.pdf')
    from docx import Document
    doc = Document()
    doc.add_heading('Structured reading fixture', 0)
    doc.add_paragraph('The human selects a source passage.')
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = 'A table cell'
    doc.save(root/'fixture.docx')
    mcp_port, viewer_port = port(), port()
    (root/'mcp-token').write_text('isolated-reading-fixture-token')
    processes = []
    logs = []
    def spawn(name, command, cwd=repo):
        log=(root/(name+'.log')).open('w'); logs.append(log)
        child=subprocess.Popen(command,cwd=cwd,env=env,stdout=log,stderr=subprocess.STDOUT)
        processes.append(child)
        return child
    try:
        spawn('extractor',[sys.executable,'-m','extraction.server'],repo/'services')
        deadline=time.monotonic()+30
        while not (root/'extractor.sock').exists():
            if time.monotonic()>deadline: raise RuntimeError('extractor did not bind')
            time.sleep(.1)
        daemon_args=[str(hades),'daemon','--database','reading_e2e','--socket',str(root/'hades.sock'),
                     '--mcp-bind',f'127.0.0.1:{mcp_port}','--mcp-token-file',str(root/'mcp-token'),
                     '--mcp-guided-reading','--mcp-dbs','reading_e2e','--mcp-db-prefix','reading_','--mcp-ingest-root',str(root)]
        spawn('daemon',daemon_args)
        deadline=time.monotonic()+30
        while not (root/'hades.sock').exists():
            if time.monotonic()>deadline: raise RuntimeError('daemon did not bind')
            time.sleep(.1)
        spawn('viewer',[str(viewer),'serve','--db','reading_e2e','--hades-bin',str(hades),
                        '--reading-db','reading_e2e','--reading-socket',str(root/'hades.sock'),
                        '--bind',f'127.0.0.1:{viewer_port}'])
        url=f'http://127.0.0.1:{viewer_port}'
        deadline=time.monotonic()+30
        while True:
            try:
                with urllib.request.urlopen(url+'/api/reading/config',timeout=1) as response:
                    assert json.load(response)['enabled']
                break
            except Exception:
                if time.monotonic()>deadline: raise
                time.sleep(.1)
        browser_env=dict(env,PLAYWRIGHT_MODULE=str(a.playwright),CHROMIUM_PATH=str(a.chromium))
        subprocess.run([str(a.node),str(repo/'scripts/test_guided_reading.cjs'),str(root),url,
                        f'http://127.0.0.1:{mcp_port}/mcp'],env=browser_env,check=True,timeout=600)
    finally:
        for child in reversed(processes):
            child.terminate()
            try: child.wait(timeout=15)
            except subprocess.TimeoutExpired:
                child.kill(); child.wait()
        for log in logs: log.close()


if __name__ == '__main__':
    main()
